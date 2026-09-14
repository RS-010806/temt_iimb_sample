#!/usr/bin/env python3
"""Build the five-page editable case study using the managed document runtime.

The evidence manifest is deliberately separate from layout. Final generation
requires actual screenshots, live URLs and completed build/API checks.
Preset: standard_business_brief. Header pattern: memo_masthead.
User-required formal overrides: Times New Roman and black for all text;
18pt title, 14pt page headings, 11pt body, 10pt tables, 9.5pt captions.
Screenshots retain their original color and decoded pixels.
"""
from __future__ import annotations

import argparse
import json
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
INK = "000000"
ORANGE = "000000"
IVORY = "FFFFFF"
MUTED = "000000"
LINE = "000000"
BODY = "Times New Roman"
DISPLAY = BODY
MONO = BODY
WIDTH = 6.5


def font(style, family, size, color=INK, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{name}"), family)
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{name}"), None)
    for name in ("spacing", "iCs"):
        for item in rpr.findall(qn(f"w:{name}")):
            rpr.remove(item)


def spacing(style, before=0, after=6, line=1.10):
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line
    style.paragraph_format.left_indent = Inches(0)
    style.paragraph_format.right_indent = Inches(0)
    style.paragraph_format.first_line_indent = Inches(0)
    ppr = style.element.get_or_add_pPr()
    for name in ("pBdr", "numPr", "contextualSpacing"):
        for item in ppr.findall(qn(f"w:{name}")):
            ppr.remove(item)


def style(doc, name, family=BODY, size=11, color=INK, bold=False, before=0, after=6, line=1.10):
    result = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    font(result, family, size, color, bold)
    spacing(result, before, after, line)
    return result


def shade(paragraph, color=IVORY):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


def field(paragraph, instruction):
    run = paragraph.add_run()
    field_el = OxmlElement("w:fldSimple")
    field_el.set(qn("w:instr"), instruction)
    run._r.addnext(field_el)


def link(paragraph, label, url):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ORANGE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def body(doc, text, name="Normal"):
    return doc.add_paragraph(text, style=name)


def label_para(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(label + " ").bold = True
    p.add_run(text)
    return p


def word_image(path):
    """Add a JFIF marker in memory for JPEG captures lacking format metadata.

    python-docx identifies JPEGs by JFIF/Exif markers. The browser can emit
    valid progressive JPEGs without either. Preserve the compressed image
    stream and assert identical decoded pixels before embedding the copy.
    """
    raw = path.read_bytes()
    if raw[:2] == b"\xff\xd8" and raw[6:11] not in (b"JFIF\x00", b"Exif\x00"):
        marker = bytes.fromhex("ffe000104a46494600010100000100010000")
        normalized = raw[:2] + marker + raw[2:]
        with Image.open(BytesIO(raw)) as original, Image.open(BytesIO(normalized)) as embedded:
            if original.size != embedded.size or original.convert("RGB").tobytes() != embedded.convert("RGB").tobytes():
                raise ValueError(f"JPEG metadata normalization changed pixels: {path}")
        return BytesIO(normalized)
    return str(path)


def screenshot(doc, path, width=WIDTH, max_height=None, alt=""):
    p = doc.add_paragraph(style="Figure")
    with Image.open(path) as img:
        ratio = img.height / img.width
    if max_height and width * ratio > max_height:
        width = max_height / ratio
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(word_image(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    return p


def screenshot_pair(doc, items, max_height):
    """Place modest-width inline captures together without layout tables."""
    p = doc.add_paragraph(style="Figure")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, (path, max_width, alt) in enumerate(items):
        with Image.open(path) as img:
            ratio = img.height / img.width
        width = min(max_width, max_height / ratio)
        shape = p.add_run().add_picture(word_image(path), width=Inches(width))
        shape._inline.docPr.set("descr", alt)
        if index < len(items) - 1:
            p.add_run("  ")
    return p


def caption(doc, number, text):
    p = doc.add_paragraph(style="Caption")
    p.add_run(f"Figure {number}. ").bold = True
    p.add_run(text)
    return p


def new_page(doc):
    p = doc.add_paragraph(style="Page Break")
    p.add_run().add_break(WD_BREAK.PAGE)


def table(doc, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(widths))
    tbl.autofit = False
    tblpr = tbl._tbl.tblPr
    width_el = tblpr.find(qn("w:tblW"))
    width_el.set(qn("w:w"), str(sum(widths)))
    width_el.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tblpr.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for edge, amount in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), str(amount))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    tblpr.append(margins)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), LINE)
        borders.append(item)
    tblpr.append(borders)
    for grid, value in zip(tbl._tbl.tblGrid.gridCol_lst, widths):
        grid.set(qn("w:w"), str(value))
    for i, values in enumerate(rows):
        row = tbl.rows[0] if i == 0 else tbl.add_row()
        for column_index, (cell, value, colwidth) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(colwidth / 1440)
            cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), str(colwidth))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            if rows[0][column_index] in {"Factor", "Desktop", "Mobile"}:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(value).bold = i == 0
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), IVORY)
                cell._tc.get_or_add_tcPr().append(shd)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
    return tbl


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    for item in doc.styles:
        if hasattr(item, "font"):
            item.font.name = BODY
            item.font.color.rgb = RGBColor.from_string(INK)
    style(doc, "Normal", BODY, 11, after=6, line=1.10)
    style(doc, "Title", BODY, 18, bold=True, after=8, line=1.05)
    style(doc, "Subtitle", BODY, 11, after=4, line=1.10)
    style(doc, "Heading 1", BODY, 14, bold=True, before=10, after=6, line=1.10)
    style(doc, "Heading 2", BODY, 12, bold=True, before=8, after=4, line=1.10)
    style(doc, "Heading 3", BODY, 11, bold=True, before=6, after=4, line=1.10)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
    page_heading = style(doc, "Page Heading", BODY, 14, bold=True, before=0, after=8, line=1.10)
    page_heading.base_style = doc.styles["Heading 1"]
    page_heading.paragraph_format.keep_with_next = True
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    page_heading.element.get_or_add_pPr().append(outline)
    style(doc, "Lead", BODY, 11, after=8, line=1.10)
    style(doc, "Figure", BODY, 1, after=0, line=1.0).paragraph_format.keep_with_next = True
    style(doc, "Caption", BODY, 9.5, before=5, after=8, line=1.10)
    style(doc, "Table Text", BODY, 10, after=0, line=1.10)
    style(doc, "Table Note", BODY, 9.5, before=5, after=4, line=1.10)
    style(doc, "After Table", BODY, 11, before=6, after=6, line=1.10)
    style(doc, "Small", BODY, 9.5, after=5, line=1.10)
    style(doc, "Link", BODY, 10, after=4, line=1.10)
    style(doc, "Page Break", BODY, 1, after=0, line=1.0)
    style(doc, "Header", BODY, 9, after=0, line=1.0)
    style(doc, "Footer", BODY, 9, after=0, line=1.0)
    for name in ("Header", "Footer"):
        doc.styles[name].paragraph_format.tab_stops.clear_all()
    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Header"]
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    hp.add_run("TEMT Technical Submission\tIIM Bangalore")
    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Footer"]
    fp.paragraph_format.tab_stops.clear_all()
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    fp.add_run("14 September 2026\tPage ")
    field(fp, "PAGE")
    fp.add_run(" of 5")
    doc.core_properties.title = "TEMT Enterprise Product Preview: Technical Submission"
    doc.core_properties.subject = "Post-meeting submission to Pratham Agarwal, IIM Bangalore"
    doc.core_properties.author = "TEMT project"
    doc.core_properties.keywords = "TEMT, technical submission, freight emissions, implementation, verification"


def enforce_formal_text(doc):
    """Make all Word text, including fields and hyperlinks, TNR and black."""
    roots = [doc.element, doc.styles.element, doc.sections[0].header._element, doc.sections[0].footer._element]
    for root in roots:
        for rpr in root.iter(qn("w:rPr")):
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.insert(0, rf)
            rf.attrib.clear()
            for name in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn(f"w:{name}"), BODY)
            color = rpr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                rpr.append(color)
            color.attrib.clear()
            color.set(qn("w:val"), INK)


def build(e):
    doc = Document()
    setup(doc)
    site = e["site_url"].rstrip("/")
    captures = {k: ROOT / v for k, v in e["screenshots"].items()}

    # 1. Formal submission and scope, supported by unchanged interface captures.
    doc.add_paragraph("TEMT Enterprise Product Preview\nTechnical Submission", style="Title")
    body(doc, "Prepared for Pratham Agarwal, IIM Bangalore", "Subtitle")
    body(doc, "Post-meeting submission | 14 September 2026 | Application release " + e["commit"], "Small")
    body(doc, "This submission records the implemented website and freight-emissions demonstration, the reasons for the principal design decisions, and the evidence used to verify the work. The deployed application and editable source are available for review.")
    screenshot_pair(doc, [
        (captures["hero"], 4.4, "Desktop production build showing the Compare chapter and its calculated 16.71 tonne emissions difference."),
        (captures["mobile"], 1.1, "Mobile production build with the estimator and input controls before the illustrative freight scene."),
    ], max_height=3.06)
    caption(doc, 1, "Desktop and mobile views of production build 81f7a0b. The desktop view supports scenario comparison; the mobile view presents inputs before the illustrative scene.")
    doc.add_paragraph("Implemented scope", style="Heading 2")
    body(doc, "The delivery comprises a responsive landing page, immediate emissions estimator, NIFTY 500 company lookup, editable shipment workspace, road/rail/ocean/air calculations, mode and monthly analytics, a validation demonstration, and a stateless analysis API. The workspace exports CSV, PDF and a Power BI import pack.")
    label_para(doc, "Audience and data.", "The principal audience is sustainability and supply-chain teams at India-focused listed enterprises. The official NIFTY 500 source snapshot contains 501 records, retrieved on 14 September 2026. Company names are real; shipment activity and resulting emissions are synthetic examples, not customer data.")
    p = body(doc, "Institutional and certification references concern the original product. The ", "Small")
    link(p, "SGS statement for TEMT v1.3", "https://dpiit.freightemissions.com/certification.pdf")
    p.add_run(" does not certify this separate demonstration engine. No exclusivity or customer relationship is asserted.")
    p = doc.add_paragraph(style="Link")
    link(p, "Deployed application", site)
    p.add_run("  |  ")
    link(p, "Source repository", e["repository_url"])

    # 2. Calculation contract and numeric example.
    new_page(doc)
    doc.add_paragraph("1. Calculation method and data", style="Page Heading")
    body(doc, "The browser and server import the same TypeScript calculator, version temt-demo-1.0.0. Each accepted transport leg is evaluated using a published well-to-wheel (WTW) default factor. This includes fuel or energy supply and vehicle operation within the stated factor boundary.")
    p = body(doc, "Emissions (kg CO2e) = tonnes × kilometres × factor (kg CO2e / tonne-km)")
    p.runs[0].bold = True
    table(doc, [
        ["Profile", "Factor", "Published scope and reference"],
        ["Road: India diesel HCV", "0.0663", "30–50 t gross vehicle weight; GLEC v3.2, p. 106"],
        ["Rail: India", "0.0106", "Mixed national traction; p. 97"],
        ["Ocean: dry container", "0.01145", "Middle East/India; converted from 114.5 g/TEU-km at 10 t/TEU; p. 112"],
        ["Air: short haul", "1.363", "Distance ≤1,500 km; p. 94"],
        ["Air: long haul", "0.788", "Distance >1,500 km; p. 94"],
    ], [2980, 1060, 5320])
    p = body(doc, "All factors are kg CO2e per tonne-km. Source: ", "Table Note")
    link(p, "Smart Freight Centre, GLEC Framework v3.2 (October 2025)", "https://smartfreightcentre.org/news/13311209")
    p.add_run(". Source URLs and assumptions are retained with the factor records.")
    doc.add_paragraph("Default estimator arithmetic", style="Heading 2")
    body(doc, "For 1,000 tonnes travelling 1,000 km by road, the baseline is 1,000 × 1,000 × 0.0663 = 66,300 kg CO2e. Moving 30% of the same activity to rail gives (700 × 1,000 × 0.0663) + (300 × 1,000 × 0.0106) = 49,590 kg CO2e. The estimated difference is 16,710 kg (16.71 t), or 25.2036%.")
    body(doc, "This compares equal transport activity and distance. It does not assess route availability, terminal movements, first/last-mile distances, capacity, cost or an operationally feasible rail substitution.")
    doc.add_paragraph("Factor and record assumptions", style="Heading 2")
    body(doc, "Published defaults already include their stated adjustments: road uses a 5% distance adjustment, 68.3% load and 14.4% empty running; ocean uses 70% fleet loading and 15% distance adjustment. Air uses a 55% belly/45% freighter mix and the published 95 km adjustment. These adjustments are not applied again. No additional aviation radiative-forcing multiplier is added. India road and rail defaults are preliminary estimates.")
    body(doc, "Each row requires shipmentId, legIndex, date, subsidiary, mode, profile, tonnes and kilometres. Dates use YYYY-MM-DD. Optional origin and destination labels describe the leg. Accepted legs aggregate by mode, month and subsidiary; filters act after validation. Compensated summation and unrounded intermediate values limit avoidable numerical error. Display rounding does not alter exported calculation inputs.")

    # 3. Validation and server contract.
    new_page(doc)
    doc.add_paragraph("2. Data validation and API controls", style="Page Heading")
    screenshot(doc, captures["pipeline"], max_height=2.25, alt="Completed three-row validation example showing 1394.33 kilograms of emissions, with a control to introduce a missing load.")
    caption(doc, 2, "Production-build validation example: three accepted legs across two shipments produce 1,394.33 kg CO2e. The missing-load control is visible but inactive.")
    label_para(doc, "Validation rules.", "Required identifiers must be non-empty and bounded in length. Leg indices must be positive integers, dates must be real calendar dates, and tonnes and kilometres must be finite and positive. Mode and profile must agree; air profiles must match the distance threshold. Numerical-range and maximum-row checks apply.")
    label_para(doc, "Shipment integrity.", "An invalid leg, duplicate leg index or missing index in a shipment excludes that entire identifiable shipment. Invalid rows without an identifiable shipment are rejected individually. Introducing the missing load in the example excludes both TEMT-002 legs; the remaining TEMT-001 leg produces 1,034.28 kg CO2e.")
    label_para(doc, "CSV import.", "Imports require a .csv file, all required headers and an explicit profile for every row, with limits of 2 MiB and 1,000 legs. Import is all-or-nothing: any parse or validation failure leaves the existing dataset unchanged. Imported rows remain in the browser until the user requests server analysis.")
    table(doc, [
        ["Endpoint", "Contract"],
        ["GET /api/health", "Status, engine version and persistence: none; /health is an alias."],
        ["GET /api/factors", "Versioned factors, methodology boundary and maximum leg count."],
        ["POST /api/analyze", 'application/json; body {"rows": [...]}; returns analysis and validation results.'],
    ], [2970, 6390])
    body(doc, "The API uses Helmet headers, no-store responses, a 2 MiB JSON limit and a default 60-request/minute/IP limit on factors and analysis. Health and CORS preflight requests are outside that limit. Exact allowed origins govern browser access; requests without an Origin header remain possible. CORS is not authentication. Controlled error responses omit internal details. Shipment bodies are processed in memory without persistence or request-body logging.", "After Table")

    # 4. Architecture and interaction rationale.
    new_page(doc)
    doc.add_paragraph("3. Architecture and interface decisions", style="Page Heading")
    label_para(doc, "Deployment.", "Next.js, React and Tailwind produce a static frontend served from Render's CDN. A separate free Express service in Singapore performs requested analyses. Both use the same calculator package; no database is required for this demonstration. CSS animation, SVG scenes, Recharts and Lucide support the interface.")
    label_para(doc, "Availability behaviour.", "The landing estimator works locally without waiting for the API. During server analysis, the interface offers a manual ‘Process in browser’ action after eight seconds. The request times out after 90 seconds; errors or timeout use the local engine. Dataset edits invalidate stale responses, and the interface identifies the result's processing origin. The free API may suspend when idle.")
    screenshot_pair(doc, [
        (captures["reporting"], 2.75, "Landing monthly analytics for the synthetic FMCG sample, totalling 453.6 tonnes of emissions."),
        (captures["demo"], 2.75, "Enterprise workspace for the same synthetic FMCG sample, including 24 shipments and 32 legs."),
    ], max_height=1.92)
    caption(doc, 3, "Landing reporting from production build 81f7a0b, and the retained public workspace capture. The visible FMCG dataset is unchanged: 24 shipments, 32 legs and 453.64 t CO2e.")
    label_para(doc, "Evaluation sequence.", "A visitor first changes freight weight, distance and rail share, then opens the workspace with those three inputs preserved in the URL. Sector and transport-mode examples expose calculated activity. The reporting preview switches between mode and month views and provides exact chart data. This sequence connects the initial estimate to inspectable records.")
    label_para(doc, "Interaction and accessibility.", "Connect, Measure, Compare and Report chapters provide context without blocking the estimator. Playback has pause and manual controls, pauses off-screen or in a hidden tab, and uses manual playback for reduced-motion or supported data-saving preferences. Mobile layouts prioritise the result and native input controls. These are design decisions, not evidence of conversion improvement.")
    label_para(doc, "Reporting.", "The current filtered view can be exported as CSV, a PDF report or a Power BI import pack. Outputs preserve factor/version context, relevant inputs and validation information. CSV formula-injection protection is tested. The Power BI pack supports importing files; it does not establish a live Power BI connection.")

    # 5. Observed evidence, handoff, and explicitly prospective work.
    new_page(doc)
    doc.add_paragraph("4. Verification, delivery and next steps", style="Page Heading")
    p = body(doc, f"Release {e['commit']} passed {e['tests_passed']} automated tests, TypeScript checks, source-integrity checks and production builds. Tests cover calculator validation, imports/exports, API controls and estimator-to-workspace parity. The ")
    link(p, "successful CI run", e["ci_url"])
    p.add_run(" records the validated application release.")
    table(doc, [
        ["Public-site Lighthouse 13.4.1", "Desktop", "Mobile"],
        ["Performance score", "100", "91"],
        ["Accessibility / best practices / SEO", "100 / 100 / 100", "100 / 100 / 100"],
        ["First contentful paint", "0.3 s", "2.0 s"],
        ["Largest contentful paint", "0.6 s", "2.9 s"],
        ["Total blocking time", "0 ms", "50 ms (54 ms raw)"],
        ["Cumulative layout shift", "0", "0"],
    ], [4860, 2160, 2340])
    body(doc, "Public audits ran on 14 September 2026 at 03:08 IST. Mobile used simulated 4× CPU slowdown and 150 ms RTT; desktop used 1× and 40 ms RTT. Performance floors of 95 desktop and 90 mobile were met. These are laboratory observations on a shared machine, not field Core Web Vitals or an accessibility certification.", "Table Note")
    label_para(doc, "API observation.", "The one-leg default scenario returned 66,300 kg CO2e with exact browser/server parity. The observed engine duration was 1.82 ms and HTTP duration, including response transfer, was 112.83 ms. An initial health request took 23.25 s with unknown prior idle time. These individual observations do not establish latency guarantees or controlled cold-start performance.")
    doc.add_paragraph("Review material", style="Heading 2")
    for labels in [
        [("Live application", site), ("Analysis workspace", site + "/demo/"), ("Source repository", e["repository_url"])],
        [("Sample shipment CSV", urljoin(site + "/", e["sample_csv_path"])), ("Sample PDF report", urljoin(site + "/", e["sample_pdf_path"])), ("Power BI import pack", urljoin(site + "/", e["power_bi_path"]))],
    ]:
        p = doc.add_paragraph(style="Link")
        for index, (label, url) in enumerate(labels):
            if index:
                p.add_run("  |  ")
            link(p, label, url)
    body(doc, "The repository retains source references, UX rationale, test records, raw Lighthouse JSON and the API measurement file in docs/. Screenshots document the interface; four show the local production build, while the unchanged workspace capture is from the prior public release.", "Small")
    doc.add_paragraph("Proposed production work and limitations", style="Heading 2")
    body(doc, "Before enterprise use, proposed work includes SSO and role-based access, tenant isolation, governed factor updates, persistent audit trails, reviewed data retention, approved logistics integrations, monitored paid capacity and repeatable load testing. None of these controls is represented as implemented in this stateless preview.")
    body(doc, "Default-factor scenarios are not carrier measurements, a complete corporate inventory, or assured reporting. Production decisions require verified activity data and reviewed methodological boundaries. Partner references, company names and original-product certification do not establish integration, endorsement or certification of this demonstration.")
    p = doc.add_paragraph(style="Small")
    link(p, "IIMB: DPIIT adoption and ULIP", "https://www.iimb.ac.in/node/14281")
    p.add_run("  |  ")
    link(p, "NSE: NIFTY 500 source snapshot", "https://www.niftyindices.com/IndexConstituent/ind_nifty500list.csv")
    enforce_formal_text(doc)
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--evidence", type=Path, default=ROOT / "docs/case-study-evidence.json")
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/deliverables/TEMT-Product-Case-Study.docx")
    args = parser.parse_args()
    e = json.loads(args.evidence.read_text())
    if e.get("status") != "ready":
        raise SystemExit("Case-study revision is awaiting final screenshots and release evidence.")
    for key in ("commit", "ci_url", "site_url", "api_url", "typecheck_passed", "source_check_passed", "production_build_passed", "live_api_verified", "live_analysis_summary", "measurements"):
        if not e.get(key):
            raise SystemExit(f"Final evidence required: {key}")
    if not isinstance(e.get("tests_passed"), int) or e["tests_passed"] < 1:
        raise SystemExit("Final passing test count required.")
    if set(e["screenshots"]) != {"hero", "mobile", "reporting", "demo", "pipeline"}:
        raise SystemExit("Five final capture roles required: hero, mobile, reporting, demo and pipeline.")
    for name, path in e["screenshots"].items():
        if not path or not (ROOT / path).is_file():
            raise SystemExit(f"Final screenshot required: {name}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc = build(e)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
